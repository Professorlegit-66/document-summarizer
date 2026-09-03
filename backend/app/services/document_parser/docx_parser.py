"""
DOCX text extraction using python-docx.
"""

import io
import zipfile
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from .exceptions import EmptyFileError, CorruptedFileError, NoTextExtractedError


def parse_docx(file_bytes: bytes) -> str:
    """
    Extract text from raw DOCX file bytes.

    Raises:
        EmptyFileError: if the file has zero bytes.
        CorruptedFileError: if the DOCX package cannot be opened.
        NoTextExtractedError: if the document opens but contains no
            meaningful text (e.g. an empty document, or only images).
    """
    if len(file_bytes) == 0:
        raise EmptyFileError("The uploaded DOCX file is empty.")

    try:
        document = Document(io.BytesIO(file_bytes))
    except (PackageNotFoundError, zipfile.BadZipFile) as exc:
        raise CorruptedFileError("The DOCX file is corrupted or not a valid Word document.") from exc

    paragraphs_text = [p.text for p in document.paragraphs if p.text.strip()]

    tables_text = []
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                tables_text.append(row_text)

    all_parts = paragraphs_text + tables_text
    full_text = "\n\n".join(all_parts)

    if not full_text.strip():
        raise NoTextExtractedError(
            "The DOCX file contains no extractable text."
        )

    return full_text